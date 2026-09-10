from docx import Document
import sys

def process_file():
    input_file = "/Users/salihacicek/Desktop/(Part B1) (HE MSCA PF) NEW.docx"
    output_file = "/Users/salihacicek/Desktop/GUNCEL_Part_B1.docx"

    try:
        doc = Document(input_file)
    except Exception as e:
        print(f"Error opening file: {e}")
        return

    wp1_text = """The integration of macroscopic 3D MRI morphology (Springbok Analytics) with localized, 2D micro-architectural data (B-mode ultrasound and SWE) presents a significant multimodal fusion challenge. Since direct voxel-to-pixel registration is geometrically intractable, we will implement a feature-based spatial alignment protocol. Both MRI and ultrasound/SWE data will be mapped to a standardized 1D anatomical coordinate system, normalized by the relative distance from the ischial tuberosity to the fibular head. This ensures precise intra- and inter-rater reliability (targeting ICC >= 0.85 and CV <= 10% for repeated measures).

Furthermore, raw sonographic and elastographic data from the field are inherently corrupted by speckle noise, acoustic shadowing, and transducer-pressure artifacts. To prevent error propagation into the downstream modelling pipeline, an autonomous, headless preprocessing script will be deployed. For B-mode US, anisotropic diffusion filtering (Perona-Malik algorithm) will be applied to iteratively smooth intra-fascicular speckle while strictly preserving the high-frequency edges of the deep and superficial aponeuroses. For SWE, elastograms will be passed through a normalized cross-correlation mask; spatial regions exhibiting a shear-wave Signal-to-Noise Ratio (SNR) of <10 dB will be autonomously flagged as invalid and excluded from the computation of the Young’s modulus, ensuring absolute data fidelity prior to predictive vectorization."""

    wp3_text = """Tree-based algorithms (Gradient-boosted trees/XGBoost) and penalised linear regressors (Elastic Net) cannot directly ingest raw 3D meshes or 2D heatmaps. Consequently, an intermediate deterministic feature extraction (vectorization) pipeline will be developed. Rather than reducing complex SWE elastograms to a single 'mean stiffness' scalar, we will extract first-order statistical moments (variance, skewness, kurtosis) and second-order textural features using Gray-Level Co-occurrence Matrices (GLCM) to mathematically quantify heterogeneous stiffness patterns within the Biceps Femoris. Similarly, 3D MRI outputs will be vectorized into spatial gradients of intramuscular adipose tissue (IMAT) and cross-sectional area (CSA) asymmetry indices.

This rigorous feature engineering will yield a high-dimensional vector space. Given the cohort size (n=300 recruited athletes), feeding this raw matrix into the XGBoost classifier would inevitably induce the curse of dimensionality and severe overfitting. To strictly enforce model generalizability across temporal and geographic validation cohorts, a two-stage dimensionality reduction architecture will be implemented. First, a collinearity filter will eliminate redundant parameters (dropping variables with a Pearson correlation coefficient |r| > 0.85). Subsequently, the Boruta algorithm will isolate statistically significant predictive features, independent of the background noise. This will condense the feature matrix to an optimal subset, maintaining a minimum Events-Per-Variable (EPV) ratio of 10:1. The final model will target an AUC/C-index >= 0.70 and Brier score <= 0.20, with full uncertainty reported via nested internal validation."""

    def add_to_paragraph(p, match_str, text_to_add):
        if match_str in p.text:
            p.add_run("\n\n" + text_to_add)
            return True
        return False

    wp1_found = False
    wp3_found = False

    # Check paragraphs
    for p in doc.paragraphs:
        if not wp1_found and add_to_paragraph(p, "O1 - Establish a harmonised", wp1_text):
            wp1_found = True
        if not wp3_found and add_to_paragraph(p, "O3 - Develop and independently evaluate", wp3_text):
            wp3_found = True

    # Check tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if not wp1_found and add_to_paragraph(p, "O1 - Establish a harmonised", wp1_text):
                        wp1_found = True
                    if not wp3_found and add_to_paragraph(p, "O3 - Develop and independently evaluate", wp3_text):
                        wp3_found = True

    try:
        doc.save(output_file)
        print(f"File saved successfully to {output_file}")
        if wp1_found: print("- WP1 text added.")
        if wp3_found: print("- WP3 text added.")
        if not wp1_found and not wp3_found:
            print("- WARNING: Did not find the markers to insert text.")
    except Exception as e:
        print(f"Error saving file: {e}")

if __name__ == "__main__":
    process_file()
